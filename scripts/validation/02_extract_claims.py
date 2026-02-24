#!/usr/bin/env python3
"""Extract figure legends and quantitative claims from manuscript DOCX."""

import argparse
import json
import re
from pathlib import Path
from docx import Document

from config import MANUSCRIPT_DOCX, EXTRACTED_DIR, COMPUTATIONAL_FIGURES, WET_LAB_FIGURES


def extract_legends(doc: Document) -> dict:
    """
    Extract figure legends from end of manuscript.

    Returns:
        dict mapping figure ID -> legend info
    """
    legends = {}
    in_legends_section = False
    current_fig = None
    current_text = []

    for para in doc.paragraphs:
        text = para.text.strip()

        if "Figures & Figure Legends" in text:
            in_legends_section = True
            continue

        if not in_legends_section:
            continue

        # Check for new figure start
        fig_match = re.match(r"^(Fig\. \d+|Extended Data Fig\.? \d+):", text)
        if fig_match:
            # Save previous figure
            if current_fig and current_text:
                legends[current_fig]["full_text"] = " ".join(current_text)

            # Start new figure
            fig_id = fig_match.group(1)
            # Normalize EDF naming
            fig_id = re.sub(r"Extended Data Fig\.? (\d+)", r"EDF \1", fig_id)

            current_fig = fig_id
            current_text = [text]

            # Determine if computational or wet lab
            fig_type = "COMPUTATIONAL" if fig_id in COMPUTATIONAL_FIGURES else "WET_LAB"

            legends[fig_id] = {
                "title": text.split(":")[1].strip()[:200] if ":" in text else text[:200],
                "type": fig_type,
                "panels": {},
                "full_text": ""
            }
        elif current_fig and text:
            current_text.append(text)

            # Extract panel descriptions (A), (B), etc.
            panel_matches = re.findall(r"\(([A-Z])\)\s*([^(]+?)(?=\([A-Z]\)|$)", text)
            for panel_letter, panel_desc in panel_matches:
                legends[current_fig]["panels"][panel_letter] = panel_desc.strip()[:300]

    # Save last figure
    if current_fig and current_text:
        legends[current_fig]["full_text"] = " ".join(current_text)

    return legends


def extract_quantitative_claims(doc: Document) -> list:
    """
    Extract quantitative claims (p-values, correlations, sample sizes) from text.

    Returns:
        list of claim dicts
    """
    claims = []

    # Patterns to match
    patterns = {
        "p_value": r"p\s*[=<>]\s*[\d.]+(?:e-?\d+)?",
        "fdr": r"FDR\s*[=<>]\s*[\d.]+",
        "correlation": r"(?:Pearson|Spearman)?\s*r\s*=\s*-?[\d.]+",
        "sample_size": r"n\s*=\s*\d+",
        "percentage": r"\d+(?:\.\d+)?%",
        "fold_change": r"\d+(?:\.\d+)?-fold",
    }

    for para_idx, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue

        # Find which figure this paragraph references
        fig_refs = re.findall(r"(?:Fig\.|Figure|Extended Data Fig\.?)\s*\d+[A-Z]?", text)

        for pattern_name, pattern in patterns.items():
            matches = re.findall(pattern, text, re.IGNORECASE)
            for match in matches:
                claims.append({
                    "type": pattern_name,
                    "value": match,
                    "context": text[:200],
                    "paragraph_idx": para_idx,
                    "figure_refs": fig_refs
                })

    return claims


def main():
    parser = argparse.ArgumentParser(description="Extract claims from manuscript")
    parser.add_argument("--legends-only", action="store_true")
    parser.add_argument("--claims-only", action="store_true")
    args = parser.parse_args()

    print(f"=== Loading manuscript: {MANUSCRIPT_DOCX.name} ===")
    doc = Document(MANUSCRIPT_DOCX)
    print(f"Total paragraphs: {len(doc.paragraphs)}")

    EXTRACTED_DIR.mkdir(parents=True, exist_ok=True)

    if not args.claims_only:
        print("\n=== Extracting figure legends ===")
        legends = extract_legends(doc)

        legends_path = EXTRACTED_DIR / "legends.json"
        with open(legends_path, "w") as f:
            json.dump(legends, f, indent=2)

        print(f"Extracted {len(legends)} figure legends")
        comp_count = sum(1 for l in legends.values() if l["type"] == "COMPUTATIONAL")
        print(f"  Computational: {comp_count}")
        print(f"  Wet lab: {len(legends) - comp_count}")
        print(f"Saved to: {legends_path}")

    if not args.legends_only:
        print("\n=== Extracting quantitative claims ===")
        claims = extract_quantitative_claims(doc)

        claims_path = EXTRACTED_DIR / "claims.json"
        with open(claims_path, "w") as f:
            json.dump(claims, f, indent=2)

        print(f"Extracted {len(claims)} quantitative claims")
        by_type = {}
        for c in claims:
            by_type[c["type"]] = by_type.get(c["type"], 0) + 1
        for t, count in sorted(by_type.items()):
            print(f"  {t}: {count}")
        print(f"Saved to: {claims_path}")


if __name__ == "__main__":
    main()
